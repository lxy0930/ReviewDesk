"""Generate a deterministic Exam Agent MAE benchmark.

The benchmark contains 100 exam papers and 100 mock student submissions.
Each paper has exactly 20 scored questions:

    8 single-choice + 4 multi-choice + 3 judge + 3 short-answer + 2 code

Public datasets provide the single-choice and code questions. The multi-choice
questions are synthesized from same-subject public options, while judge and
short-answer questions come from the local curated question bank.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import random
import re
import sys
import time
import uuid
from pathlib import Path
from typing import Any

import httpx
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Pt

if sys.stdout and hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
if sys.stderr and hasattr(sys.stderr, "reconfigure"):
    sys.stderr.reconfigure(encoding="utf-8")


ROOT = Path(__file__).resolve().parent
GENERATED_DIR = ROOT / "generated"
CACHE_DIR = ROOT / "source_cache"
EXAM_DIR = GENERATED_DIR / "exam_papers"
ANSWER_KEY_DIR = GENERATED_DIR / "answer_keys"
SUBMISSION_DIR = GENERATED_DIR / "submissions"
GROUND_TRUTH_DIR = GENERATED_DIR / "ground_truth"

DATASET_SERVER = "https://datasets-server.huggingface.co/rows"
UUID_NAMESPACE = uuid.UUID("3e94694b-8ef0-40d4-bd42-4ccb13b5bb9a")

SINGLE_SOURCES = [
    ("ceval/ceval-exam", "college_programming"),
    ("ceval/ceval-exam", "computer_architecture"),
    ("ceval/ceval-exam", "computer_network"),
    ("ceval/ceval-exam", "operating_system"),
    ("ceval/ceval-exam", "discrete_mathematics"),
    ("cais/mmlu", "college_computer_science"),
    ("cais/mmlu", "computer_security"),
    ("cais/mmlu", "high_school_computer_science"),
    ("cais/mmlu", "machine_learning"),
]

CODE_SOURCES = [
    ("openai/openai_humaneval", "openai_humaneval", "test"),
    ("google-research-datasets/mbpp", "sanitized", "test"),
]

TYPE_LABELS = {
    "single_choice": "单选",
    "multi_choice": "多选",
    "judge": "判断",
    "short_answer": "简答",
    "code": "代码",
}

PROFILE_RULES = {
    "strong": {
        "objective_accuracy": 0.92,
        "code_score_rate": 1.0,
    },
    "mixed": {
        "objective_accuracy": 0.58,
        "code_score_rate": 0.5,
    },
    "weak": {
        "objective_accuracy": 0.22,
        "code_score_rate": 0.0,
    },
}


def stable_id(*parts: str) -> uuid.UUID:
    return uuid.uuid5(UUID_NAMESPACE, "|".join(parts))


def deterministic_order(seed: int, value: str) -> str:
    return hashlib.sha256(f"{seed}:{value}".encode("utf-8")).hexdigest()


def normalize_text(value: str) -> str:
    return re.sub(r"\s+", " ", value or "").strip()


def fetch_rows(
    client: httpx.Client,
    dataset: str,
    config: str,
    split: str,
    limit: int,
    force_download: bool,
) -> list[dict[str, Any]]:
    cache_name = (
        dataset.replace("/", "__")
        + "__"
        + config.replace("/", "__")
        + "__"
        + split
        + ".json"
    )
    cache_path = CACHE_DIR / cache_name

    if cache_path.exists() and not force_download:
        cached = json.loads(cache_path.read_text(encoding="utf-8"))
        if len(cached) >= limit:
            return cached[:limit]

    rows: list[dict[str, Any]] = []
    offset = 0
    while len(rows) < limit:
        page_size = min(100, limit - len(rows))
        response = None
        last_error: Exception | None = None
        for attempt in range(3):
            try:
                response = client.get(
                    DATASET_SERVER,
                    params={
                        "dataset": dataset,
                        "config": config,
                        "split": split,
                        "offset": offset,
                        "length": page_size,
                    },
                )
                break
            except httpx.TransportError as error:
                last_error = error
                if attempt < 2:
                    time.sleep(2**attempt)
        if response is None:
            raise RuntimeError(
                f"下载数据集失败: {dataset}:{config}: {last_error}"
            )
        response.raise_for_status()
        payload = response.json()
        page = payload.get("rows", [])
        if not page:
            break
        rows.extend(page)
        offset += len(page)

    CACHE_DIR.mkdir(parents=True, exist_ok=True)
    cache_path.write_text(
        json.dumps(rows, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    return rows[:limit]


def normalize_single_choice(
    source_dataset: str,
    config: str,
    rows: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    normalized: list[dict[str, Any]] = []
    seen: set[str] = set()

    for row_wrapper in rows:
        row = row_wrapper.get("row", row_wrapper)
        if source_dataset == "ceval/ceval-exam":
            question = normalize_text(str(row.get("question", "")))
            option_map = {
                letter: normalize_text(str(row.get(letter, "")))
                for letter in ("A", "B", "C", "D")
            }
            answer = str(row.get("answer", "")).strip().upper()
            if answer not in option_map:
                continue
        else:
            question = normalize_text(str(row.get("question", "")))
            choices = row.get("choices") or []
            if len(choices) != 4:
                continue
            letters = ("A", "B", "C", "D")
            option_map = {
                letter: normalize_text(str(choices[index]))
                for index, letter in enumerate(letters)
            }
            answer_index = row.get("answer")
            if not isinstance(answer_index, int) or not 0 <= answer_index < 4:
                continue
            answer = letters[answer_index]

        options = [option_map[letter] for letter in ("A", "B", "C", "D")]
        if not question or any(not option for option in options):
            continue

        signature = hashlib.sha1(
            json.dumps(
                [question, options, answer],
                ensure_ascii=False,
                sort_keys=True,
            ).encode("utf-8")
        ).hexdigest()
        if signature in seen:
            continue
        seen.add(signature)

        normalized.append(
            {
                "source_id": f"{source_dataset}:{config}:{row_wrapper.get('row_idx')}",
                "source_dataset": source_dataset,
                "subject": config,
                "question": question,
                "options": options,
                "correct_letter": answer,
                "correct_text": option_map[answer],
                "incorrect_options": [
                    option_map[letter]
                    for letter in ("A", "B", "C", "D")
                    if letter != answer
                ],
            }
        )

    return normalized


def build_multi_choice_pool(
    single_pool: list[dict[str, Any]],
    target_count: int,
    rng: random.Random,
) -> list[dict[str, Any]]:
    by_subject: dict[str, list[dict[str, Any]]] = {}
    for item in single_pool:
        by_subject.setdefault(item["subject"], []).append(item)

    subjects = list(by_subject)
    if not subjects:
        return []

    pool: list[dict[str, Any]] = []
    seen: set[str] = set()
    attempts = 0

    while len(pool) < target_count and attempts < target_count * 100:
        attempts += 1
        subject = rng.choice(subjects)
        items = by_subject[subject]
        if len(items) < 4:
            continue

        chosen = rng.sample(items, 4)
        correct_items = chosen[:2]
        wrong_items = chosen[2:]

        correct_texts = [item["correct_text"] for item in correct_items]
        wrong_texts: list[str] = []
        for item in wrong_items:
            candidates = [
                option
                for option in item["incorrect_options"]
                if option not in correct_texts and option not in wrong_texts
            ]
            if not candidates:
                break
            wrong_texts.append(rng.choice(candidates))
        if len(wrong_texts) != 2:
            continue

        option_entries = [
            {"text": text, "correct": True}
            for text in correct_texts
        ] + [
            {"text": text, "correct": False}
            for text in wrong_texts
        ]
        if len({entry["text"] for entry in option_entries}) != 4:
            continue

        rng.shuffle(option_entries)
        answer_letters = [
            "ABCD"[index]
            for index, entry in enumerate(option_entries)
            if entry["correct"]
        ]
        signature = hashlib.sha1(
            json.dumps(option_entries, ensure_ascii=False).encode("utf-8")
        ).hexdigest()
        if signature in seen:
            continue
        seen.add(signature)

        pool.append(
            {
                "source_id": f"synthetic-multi:{signature[:16]}",
                "source_dataset": "synthetic-from-ceval-mmlu",
                "subject": subject,
                "question": f"关于 {subject}，以下哪些说法是正确的？",
                "options": [entry["text"] for entry in option_entries],
                "correct_letter": "".join(answer_letters),
                "correct_text": "；".join(
                    entry["text"]
                    for entry in option_entries
                    if entry["correct"]
                ),
            }
        )

    return pool


def normalize_code_questions(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    normalized: list[dict[str, Any]] = []
    seen: set[str] = set()

    for wrapper in rows:
        row = wrapper.get("row", wrapper)
        task_id = str(row.get("task_id", "")).strip()
        # Code prompts contain newlines and indentation. Trimming is safe, but
        # whitespace normalization would collapse the code into one line.
        question = str(row.get("prompt", "")).strip()
        if not question:
            continue

        if "canonical_solution" in row:
            prompt = str(row.get("prompt", ""))
            solution = str(row.get("canonical_solution", ""))
            reference_code = (prompt + solution).strip()
            source_dataset = "openai/openai_humaneval"
        else:
            reference_code = str(row.get("code", "")).strip()
            source_dataset = "google-research-datasets/mbpp"

        if not reference_code:
            continue
        signature = hashlib.sha1(
            f"{source_dataset}:{task_id}:{question}".encode("utf-8")
        ).hexdigest()
        if signature in seen:
            continue
        seen.add(signature)

        normalized.append(
            {
                "source_id": f"{source_dataset}:{task_id}",
                "source_dataset": source_dataset,
                "subject": "programming",
                "question": question,
                "reference_code": reference_code,
            }
        )

    return normalized


def choose_wrong_letter(
    correct_letter: str,
    rng: random.Random,
) -> str:
    choices = [letter for letter in "ABCD" if letter != correct_letter]
    return rng.choice(choices)


def choose_multi_answer(
    correct_letters: str,
    profile: str,
    rng: random.Random,
) -> str:
    correct = set(correct_letters)
    if profile == "strong":
        return "".join(sorted(correct))

    if profile == "mixed":
        keep = set(sorted(correct)[:1])
        if rng.random() < 0.5:
            wrong = rng.choice([letter for letter in "ABCD" if letter not in correct])
            keep.add(wrong)
        return "".join(sorted(keep))

    subset = {letter for letter in "ABCD" if rng.random() < 0.35}
    if not subset:
        subset = {rng.choice("ABCD")}
    return "".join(sorted(subset))


def build_short_answer(
    bank_item: dict[str, Any],
    profile: str,
) -> tuple[str, int]:
    points = bank_item["points"]
    reference_answer = bank_item["reference_answer"].strip()
    answer_parts = [
        part.strip()
        for part in re.split(r"[；;]", reference_answer)
        if part.strip()
    ]

    if profile == "strong":
        return reference_answer, sum(point["score"] for point in points)

    if profile == "mixed":
        partial_answer = (
            answer_parts[0]
            if answer_parts
            else reference_answer
        )
        return partial_answer, points[0]["score"]

    if profile == "weak":
        return "不了解，暂时无法作答。", 0

    raise ValueError(f"未知 student profile: {profile}")


def build_code_answer(
    reference_code: str,
    profile: str,
    max_score: int,
) -> tuple[str, int]:
    if profile == "strong":
        return reference_code, max_score

    if profile == "mixed":
        return "# 只完成了部分逻辑\npass", round(max_score * 0.5)

    return "# 未完成\npass", 0


def question_uid(paper_id: str, question_no: int) -> str:
    return str(stable_id("question", paper_id, str(question_no)))


def build_paper(
    paper_index: int,
    rng: random.Random,
    single_pool: list[dict[str, Any]],
    multi_pool: list[dict[str, Any]],
    code_pool: list[dict[str, Any]],
    curated: dict[str, Any],
) -> dict[str, Any]:
    paper_no = paper_index + 1
    paper_id = f"exam_{paper_no:03d}"
    case_id = f"case_{paper_no:03d}"
    profile = ("strong", "mixed", "weak")[paper_index % 3]

    exam_uuid = str(stable_id("exam", paper_id))

    single_items = rng.sample(single_pool, 8)
    multi_items = rng.sample(multi_pool, 4)

    judge_offset = paper_index % len(curated["judge"])
    judge_items = [
        curated["judge"][(judge_offset + index) % len(curated["judge"])]
        for index in range(3)
    ]

    short_offset = paper_index % len(curated["short_answer"])
    short_items = [
        curated["short_answer"][
            (short_offset + index) % len(curated["short_answer"])
        ]
        for index in range(3)
    ]

    code_start = (paper_index * 2) % max(1, len(code_pool) - 1)
    code_items = [
        code_pool[(code_start + index) % len(code_pool)]
        for index in range(2)
    ]

    questions: list[dict[str, Any]] = []

    for item in single_items:
        question_no = len(questions) + 1
        correct = item["correct_letter"]
        if rng.random() < PROFILE_RULES[profile]["objective_accuracy"]:
            student_letter = correct
        else:
            student_letter = choose_wrong_letter(correct, rng)

        max_score = 3
        human_score = max_score if student_letter == correct else 0
        questions.append(
            {
                "question_id": question_uid(paper_id, question_no),
                "question_no": question_no,
                "question_type": "single_choice",
                "max_score": max_score,
                "content": item["question"],
                "options": item["options"],
                "answer": correct,
                "answer_text": item["correct_text"],
                "student_answer": student_letter,
                "human_score": human_score,
                "knowledge_tag": item["subject"],
                "source_id": item["source_id"],
                "source_dataset": item["source_dataset"],
            }
        )

    for item in multi_items:
        question_no = len(questions) + 1
        correct = item["correct_letter"]
        student_answer = choose_multi_answer(correct, profile, rng)
        max_score = 4
        human_score = max_score if student_answer == correct else 0
        questions.append(
            {
                "question_id": question_uid(paper_id, question_no),
                "question_no": question_no,
                "question_type": "multi_choice",
                "max_score": max_score,
                "content": item["question"],
                "options": item["options"],
                "answer": correct,
                "answer_text": item["correct_text"],
                "student_answer": student_answer,
                "human_score": human_score,
                "knowledge_tag": item["subject"],
                "source_id": item["source_id"],
                "source_dataset": item["source_dataset"],
            }
        )

    for item in judge_items:
        question_no = len(questions) + 1
        correct = "对" if item["answer"] else "错"
        if rng.random() < PROFILE_RULES[profile]["objective_accuracy"]:
            student_answer = correct
        else:
            student_answer = "错" if correct == "对" else "对"

        max_score = 2
        human_score = max_score if student_answer == correct else 0
        questions.append(
            {
                "question_id": question_uid(paper_id, question_no),
                "question_no": question_no,
                "question_type": "judge",
                "max_score": max_score,
                "content": item["question"],
                "options": [],
                "answer": correct,
                "answer_text": item["explanation"],
                "student_answer": student_answer,
                "human_score": human_score,
                "knowledge_tag": "计算机基础",
                "source_id": "curated:judge",
                "source_dataset": "curated-local",
            }
        )

    for item in short_items:
        question_no = len(questions) + 1
        student_answer, human_score = build_short_answer(item, profile)
        questions.append(
            {
                "question_id": question_uid(paper_id, question_no),
                "question_no": question_no,
                "question_type": "short_answer",
                "max_score": 10,
                "content": item["question"],
                "options": [],
                "answer": item["reference_answer"],
                "answer_text": item["reference_answer"],
                "scoring_points": item["points"],
                "student_answer": student_answer,
                "human_score": human_score,
                "knowledge_tag": "计算机基础与软件工程",
                "source_id": "curated:short_answer",
                "source_dataset": "curated-local",
            }
        )

    for item in code_items:
        question_no = len(questions) + 1
        student_answer, human_score = build_code_answer(
            item["reference_code"],
            profile,
            12,
        )
        questions.append(
            {
                "question_id": question_uid(paper_id, question_no),
                "question_no": question_no,
                "question_type": "code",
                "max_score": 12,
                "content": item["question"],
                "options": [],
                "answer": item["reference_code"],
                "answer_text": item["reference_code"],
                "student_answer": student_answer,
                "human_score": human_score,
                "knowledge_tag": "编程与算法",
                "source_id": item["source_id"],
                "source_dataset": item["source_dataset"],
            }
        )

    return {
        "paper_id": paper_id,
        "exam_id": exam_uuid,
        "case_id": case_id,
        "profile": profile,
        "questions": questions,
    }


def write_exam_docx(paper: dict[str, Any], output_path: Path) -> None:
    document = Document()
    document.add_heading(f"IT 综合能力测试 {paper['paper_id']}", level=0)
    document.add_paragraph(
        "本试卷共 20 题，满分 100 分。请在第 1 题至第 20 题下作答。"
    )

    for question in paper["questions"]:
        document.add_paragraph(
            f"第{question['question_no']}题"
            f"（{TYPE_LABELS[question['question_type']]}，"
            f"{question['max_score']}分）"
        )
        if question["question_type"] == "code":
            add_code_block(document, question["content"])
        else:
            document.add_paragraph(question["content"])

        if question["options"]:
            document.add_paragraph("选项：")
            for index, option in enumerate(question["options"]):
                paragraph = document.add_paragraph()
                paragraph.add_run(f"{'ABCD'[index]}. ").bold = True
                paragraph.add_run(option)

        if question["question_type"] == "short_answer":
            document.add_paragraph("作答区：")
        elif question["question_type"] == "code":
            document.add_paragraph("请在代码块中提交完整实现。")

    document.save(output_path)


def add_code_block(document: Document, code: str) -> None:
    """Write code as separate monospaced paragraphs with preserved indentation."""
    for line in code.splitlines() or [""]:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        run = paragraph.add_run(line if line else " ")
        run.font.name = "Consolas"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        run.font.size = Pt(9)


def write_answer_key(
    paper: dict[str, Any],
    output_path: Path,
) -> None:
    payload = {
        "paper_id": paper["paper_id"],
        "exam_id": paper["exam_id"],
        "questions": [
            {
                "question_id": question["question_id"],
                "question_no": question["question_no"],
                "question_type": question["question_type"],
                "max_score": question["max_score"],
                "answer": question["answer"],
                "answer_text": question["answer_text"],
                "scoring_points": question.get("scoring_points", []),
                "source_id": question["source_id"],
                "source_dataset": question["source_dataset"],
            }
            for question in paper["questions"]
        ],
    }
    output_path.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )


def write_student_docx(
    paper: dict[str, Any],
    output_path: Path,
) -> None:
    document = Document()
    document.add_heading(
        f"提交答卷 {paper['case_id']} / {paper['paper_id']}",
        level=1,
    )

    for question in paper["questions"]:
        document.add_paragraph(f"第{question['question_no']}题")
        if question["question_type"] == "code":
            add_code_block(
                document,
                "```python\n"
                + str(question["student_answer"])
                + "\n```",
            )
        else:
            document.add_paragraph(f"答：{question['student_answer']}")

    document.save(output_path)


def write_csv(
    path: Path,
    rows: list[dict[str, Any]],
    fieldnames: list[str],
) -> None:
    with path.open("w", encoding="utf-8-sig", newline="") as file:
        writer = csv.DictWriter(file, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def generate(args: argparse.Namespace) -> None:
    if args.papers <= 0:
        raise SystemExit("--papers 必须大于 0")

    rng = random.Random(args.seed)

    for directory in (
        GENERATED_DIR,
        CACHE_DIR,
        EXAM_DIR,
        ANSWER_KEY_DIR,
        SUBMISSION_DIR,
        GROUND_TRUTH_DIR,
    ):
        directory.mkdir(parents=True, exist_ok=True)

    curated = json.loads(
        (ROOT / "curated_questions.json").read_text(encoding="utf-8")
    )

    single_pool: list[dict[str, Any]] = []
    source_counts: dict[str, int] = {}

    with httpx.Client(timeout=60.0) as client:
        for dataset, config in SINGLE_SOURCES:
            rows = fetch_rows(
                client,
                dataset,
                config,
                "test",
                100,
                args.force_download,
            )
            normalized = normalize_single_choice(dataset, config, rows)
            single_pool.extend(normalized)
            source_counts[f"{dataset}:{config}"] = len(normalized)

        code_pool: list[dict[str, Any]] = []
        for dataset, config, split in CODE_SOURCES:
            rows = fetch_rows(
                client,
                dataset,
                config,
                split,
                300,
                args.force_download,
            )
            normalized = normalize_code_questions(rows)
            code_pool.extend(normalized)
            source_counts[f"{dataset}:{config}"] = len(normalized)

    if len(single_pool) < 100:
        raise RuntimeError(
            f"单选题池不足：{len(single_pool)}，至少需要 100"
        )
    if len(code_pool) < 20:
        raise RuntimeError(
            f"代码题池不足：{len(code_pool)}，至少需要 20"
        )

    multi_pool = build_multi_choice_pool(
        single_pool,
        target_count=max(100, args.papers * 4),
        rng=rng,
    )
    if len(multi_pool) < 4:
        raise RuntimeError("多选题池不足")

    papers: list[dict[str, Any]] = []
    manifest_items: list[dict[str, Any]] = []
    human_score_rows: list[dict[str, Any]] = []

    for paper_index in range(args.papers):
        paper = build_paper(
            paper_index=paper_index,
            rng=rng,
            single_pool=single_pool,
            multi_pool=multi_pool,
            code_pool=code_pool,
            curated=curated,
        )
        papers.append(paper)

        exam_path = EXAM_DIR / f"{paper['paper_id']}.docx"
        answer_key_path = ANSWER_KEY_DIR / f"{paper['paper_id']}_answer_key.json"
        submission_path = SUBMISSION_DIR / f"{paper['case_id']}.docx"

        write_exam_docx(paper, exam_path)
        write_answer_key(paper, answer_key_path)
        write_student_docx(paper, submission_path)

        for question in paper["questions"]:
            human_score_rows.append(
                {
                    "paper_id": paper["paper_id"],
                    "case_id": paper["case_id"],
                    "question_no": question["question_no"],
                    "question_type": question["question_type"],
                    "max_score": question["max_score"],
                    "human_score": question["human_score"],
                    "source_dataset": question["source_dataset"],
                }
            )

        manifest_items.append(
            {
                "paper_id": paper["paper_id"],
                "exam_id": paper["exam_id"],
                "case_id": paper["case_id"],
                "profile": paper["profile"],
                "exam_file": str(exam_path.relative_to(GENERATED_DIR)),
                "answer_key_file": str(
                    answer_key_path.relative_to(GENERATED_DIR)
                ),
                "submission_file": str(
                    submission_path.relative_to(GENERATED_DIR)
                ),
            }
        )

        print(
            f"generated {paper['paper_id']} "
            f"profile={paper['profile']} "
            f"human_total={sum(q['human_score'] for q in paper['questions'])}"
        )

    dataset_path = GENERATED_DIR / "benchmark_dataset.json"
    dataset_path.write_text(
        json.dumps(
            {
                "seed": args.seed,
                "paper_count": len(papers),
                "questions_per_paper": 20,
                "score_total": 100,
                "papers": papers,
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )

    manifest_path = GENERATED_DIR / "manifest.json"
    manifest_path.write_text(
        json.dumps(
            {
                "seed": args.seed,
                "papers": manifest_items,
                "source_counts": source_counts,
                "single_pool_size": len(single_pool),
                "multi_pool_size": len(multi_pool),
                "code_pool_size": len(code_pool),
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )

    write_csv(
        GROUND_TRUTH_DIR / "all_human_scores.csv",
        human_score_rows,
        [
            "paper_id",
            "case_id",
            "question_no",
            "question_type",
            "max_score",
            "human_score",
            "source_dataset",
        ],
    )

    print(f"\ngenerated papers: {len(papers)}")
    print(f"generated submissions: {len(papers)}")
    print(f"single pool: {len(single_pool)}")
    print(f"multi pool: {len(multi_pool)}")
    print(f"code pool: {len(code_pool)}")
    print(f"output: {GENERATED_DIR}")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "--papers",
        type=int,
        default=100,
        help="生成试卷数量，默认 100",
    )
    parser.add_argument(
        "--seed",
        type=int,
        default=20260927,
        help="随机种子",
    )
    parser.add_argument(
        "--force-download",
        action="store_true",
        help="忽略 source_cache，重新下载公开数据集",
    )
    return parser.parse_args()


if __name__ == "__main__":
    generate(parse_args())
