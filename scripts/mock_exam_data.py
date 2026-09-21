# scripts/mock_exam_data.py
# 执行：python scripts/mock_exam_data.py
# 用途：灌入一份接近真实考试规模的 Python 基础试卷 mock 数据，
#       并生成学生作答 Word 文件。
#
# 题目来源：
#   - https://docs.python.org/3/tutorial/
#   - https://www.w3schools.com/python/
#   - https://realpython.com/python-basics/
#
# 依赖：需先运行 scripts/seed_data.py 创建 teacher01 / student01。

import asyncio
import os
import sys

if sys.stdout and hasattr(sys.stdout, "reconfigure"):
    try:
        sys.stdout.reconfigure(encoding="utf-8")
        sys.stderr.reconfigure(encoding="utf-8")
    except Exception:
        pass

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import asyncpg
from backend.config import get_settings

s = get_settings()
DB_DSN = (
    f"postgresql://{s.db_user}:{s.db_password}"
    f"@{s.db_host}:{s.db_port}/{s.db_name}"
)

EXAM_ID = "10000000-0000-0000-0000-000000000001"
STUDENT_CASES = [
    {
        "username": "student01",
        "email": "student01@eduagent.local",
        "submission_id": "40000000-0000-0000-0000-000000000001",
        "variant": "strong",
    },
    {
        "username": "student02",
        "email": "student02@eduagent.local",
        "submission_id": "40000000-0000-0000-0000-000000000002",
        "variant": "mixed",
    },
    {
        "username": "student03",
        "email": "student03@eduagent.local",
        "submission_id": "40000000-0000-0000-0000-000000000003",
        "variant": "weak",
    },
]

Q_IDS = {
    i: f"20000000-0000-0000-0000-{i:012d}"
    for i in range(1, 21)
}

SP_IDS = {
    i: f"30000000-0000-0000-0000-{i:012d}"
    for i in range(1, 7)
}

EXAM = {
    "title": "Python 基础能力 · 综合模拟测试",
    "description": "覆盖 Python 基础语法、数据类型、函数、类、文件操作和基础算法，共 20 题。",
}

QUESTIONS = [
    {
        "no": 1,
        "type": "single_choice",
        "score": 3,
        "tag": "Python基础",
        "content": "Python 中用于定义函数的关键字是？\nA. func\nB. def\nC. function\nD. define",
        "answer": "B",
    },
    {
        "no": 2,
        "type": "single_choice",
        "score": 3,
        "tag": "数据类型",
        "content": "下列哪种数据类型是不可变的？\nA. list\nB. tuple\nC. dict\nD. set",
        "answer": "B",
    },
    {
        "no": 3,
        "type": "single_choice",
        "score": 3,
        "tag": "运算符",
        "content": "表达式 2 ** 3 的结果是？\nA. 6\nB. 8\nC. 9\nD. 16",
        "answer": "B",
    },
    {
        "no": 4,
        "type": "single_choice",
        "score": 3,
        "tag": "列表",
        "content": "向列表末尾添加一个元素，应该使用哪个方法？\nA. add()\nB. append()\nC. insert()\nD. extend()",
        "answer": "B",
    },
    {
        "no": 5,
        "type": "single_choice",
        "score": 3,
        "tag": "模块",
        "content": "Python 标准库中用于处理正则表达式的模块是？\nA. re\nB. regex\nC. pyre\nD. reg",
        "answer": "A",
    },
    {
        "no": 6,
        "type": "single_choice",
        "score": 3,
        "tag": "文件操作",
        "content": "以只读方式打开文件 file.txt，正确写法是？\nA. open('file.txt', 'r')\nB. read('file.txt')\nC. open_file('file.txt')\nD. load('file.txt')",
        "answer": "A",
    },
    {
        "no": 7,
        "type": "single_choice",
        "score": 3,
        "tag": "内置函数",
        "content": "获取列表长度的内置函数是？\nA. size()\nB. length()\nC. len()\nD. count()",
        "answer": "C",
    },
    {
        "no": 8,
        "type": "single_choice",
        "score": 3,
        "tag": "运算符",
        "content": "Python 中整除运算符是？\nA. //\nB. /\nC. %\nD. **",
        "answer": "A",
    },
    {
        "no": 9,
        "type": "multi_choice",
        "score": 4,
        "tag": "数据类型",
        "content": "以下哪些属于 Python 内置数据类型？\nA. list\nB. dict\nC. array\nD. tuple",
        "answer": "ABD",
    },
    {
        "no": 10,
        "type": "multi_choice",
        "score": 4,
        "tag": "数据类型",
        "content": "以下哪些类型是可变类型？\nA. list\nB. dict\nC. tuple\nD. set",
        "answer": "ABD",
    },
    {
        "no": 11,
        "type": "multi_choice",
        "score": 4,
        "tag": "控制流",
        "content": "以下哪些关键字可以用于创建循环？\nA. for\nB. while\nC. if\nD. def",
        "answer": "AB",
    },
    {
        "no": 12,
        "type": "judge",
        "score": 2,
        "tag": "Python基础",
        "content": "Python 是编译型语言。",
        "answer": "错",
    },
    {
        "no": 13,
        "type": "judge",
        "score": 2,
        "tag": "Python基础",
        "content": "在 Python 中，None 等价于 False。",
        "answer": "错",
    },
    {
        "no": 14,
        "type": "judge",
        "score": 2,
        "tag": "集合",
        "content": "Python 的 set 允许存储重复元素。",
        "answer": "错",
    },
    {
        "no": 15,
        "type": "short_answer",
        "score": 10,
        "tag": "数据类型",
        "content": "请说明列表 list 与元组 tuple 的主要区别。",
        "answer": "列表可变，元组不可变；列表使用方括号，元组使用圆括号；列表适合增删改，元组适合不可变数据。",
        "points": [
            {"id": 1, "desc": "指出列表可变、元组不可变", "score": 5},
            {"id": 2, "desc": "指出语法或使用场景差异", "score": 5},
        ],
    },
    {
        "no": 16,
        "type": "short_answer",
        "score": 10,
        "tag": "递归算法",
        "content": "什么是递归？请说明递归的两个必要条件。",
        "answer": "递归是函数调用自身的过程；两个必要条件是基线条件和递归条件。",
        "points": [
            {"id": 3, "desc": "正确描述递归定义", "score": 5},
            {"id": 4, "desc": "正确说明基线条件和递归条件", "score": 5},
        ],
    },
    {
        "no": 17,
        "type": "short_answer",
        "score": 10,
        "tag": "面向对象",
        "content": "Python 类中的 __init__ 方法有什么作用？",
        "answer": "__init__ 是构造方法，在创建对象时自动调用，用于初始化对象属性。",
        "points": [
            {"id": 5, "desc": "说明 __init__ 是构造方法", "score": 5},
            {"id": 6, "desc": "说明用于初始化对象属性", "score": 5},
        ],
    },
    {
        "no": 18,
        "type": "code",
        "score": 12,
        "tag": "动态规划",
        "content": "编写函数 fibonacci(n)，返回斐波那契数列的第 n 项，约定 fibonacci(0)=0，fibonacci(1)=1。",
        "answer": "def fibonacci(n):\n    if n < 2:\n        return n\n    a, b = 0, 1\n    for _ in range(n - 1):\n        a, b = b, a + b\n    return b",
    },
    {
        "no": 19,
        "type": "code",
        "score": 12,
        "tag": "字符串",
        "content": "编写函数 char_count(text)，返回一个字典，统计字符串中每个字符出现的次数。",
        "answer": "def char_count(text):\n    result = {}\n    for ch in text:\n        result[ch] = result.get(ch, 0) + 1\n    return result",
    },
    {
        "no": 20,
        "type": "code",
        "score": 12,
        "tag": "字符串",
        "content": "编写函数 is_palindrome(text)，判断字符串是否为回文，忽略空格和大小写。",
        "answer": "def is_palindrome(text):\n    text = ''.join(ch.lower() for ch in text if ch.isalnum())\n    return text == text[::-1]",
    },
]


async def seed():
    """灌入试卷、20 道题、得分点和学生提交记录。"""
    conn = await asyncpg.connect(DB_DSN)
    print(f"✅ 数据库连接成功：{s.db_host}:{s.db_port}/{s.db_name}")
    try:
        teacher_id = await conn.fetchval(
            "SELECT id FROM users WHERE email = 'teacher01@eduagent.local'"
        )
        print(f"teacher01.id = {teacher_id}")

        if teacher_id is None:
            print("⚠️  未找到 teacher01，请先运行：python scripts/seed_data.py")
            return

        await conn.execute(
            """
            INSERT INTO exams (id, title, description, created_by, is_active)
            VALUES ($1, $2, $3, $4, TRUE)
            ON CONFLICT (id) DO NOTHING
            """,
            EXAM_ID,
            EXAM["title"],
            EXAM["description"],
            str(teacher_id),
        )
        print(f"📄 试卷：{EXAM['title']}  (exam_id={EXAM_ID})")

        for q in QUESTIONS:
            await conn.execute(
                """
                INSERT INTO questions (
                    id, exam_id, question_no, question_type,
                    content, correct_answer, score, knowledge_tag
                )
                VALUES ($1, $2, $3, $4, $5, $6, $7, $8)
                ON CONFLICT (id) DO UPDATE
                SET question_no = EXCLUDED.question_no,
                    question_type = EXCLUDED.question_type,
                    content = EXCLUDED.content,
                    correct_answer = EXCLUDED.correct_answer,
                    score = EXCLUDED.score,
                    knowledge_tag = EXCLUDED.knowledge_tag
                """,
                Q_IDS[q["no"]],
                EXAM_ID,
                q["no"],
                q["type"],
                q["content"],
                q["answer"],
                q["score"],
                q["tag"],
            )
            print(f"   题{q['no']} [{q['type']}] {q['score']}分 知识点={q['tag']}")

        for q in QUESTIONS:
            for sp in q.get("points", []):
                await conn.execute(
                    """
                    INSERT INTO scoring_points (
                        id, question_id, point_desc, point_score, is_active
                    )
                    VALUES ($1, $2, $3, $4, TRUE)
                    ON CONFLICT (id) DO UPDATE
                    SET question_id = EXCLUDED.question_id,
                        point_desc = EXCLUDED.point_desc,
                        point_score = EXCLUDED.point_score,
                        is_active = TRUE
                    """,
                    SP_IDS[sp["id"]],
                    Q_IDS[q["no"]],
                    sp["desc"],
                    sp["score"],
                )
                print(f"       ↳ 得分点：{sp['desc']}（{sp['score']}分）")

        for case in STUDENT_CASES:
            student_id = await conn.fetchval(
                "SELECT id FROM users WHERE email = $1",
                case["email"],
            )
            if student_id is None:
                print(f"⚠️  跳过 {case['username']}：账号不存在")
                continue

            word_minio_path = (
                f"exams/{student_id}/{case['submission_id']}.docx"
            )
            await conn.execute(
                """
                DELETE FROM exam_submissions
                WHERE exam_id = $1 AND student_id = $2
                """,
                EXAM_ID,
                str(student_id),
            )
            await conn.execute(
                """
                INSERT INTO exam_submissions (
                    id, exam_id, student_id, source, word_minio_path, status
                )
                VALUES ($1, $2, $3, 'word', $4, 'submitted')
                """,
                case["submission_id"],
                EXAM_ID,
                str(student_id),
                word_minio_path,
            )
            print(f"🧑‍🎓 学生提交：{case['username']} → {EXAM['title']}")
            print(f"   submission_id = {case['submission_id']}")
            print(f"   word_minio_path = {word_minio_path}")
            print(f"   variant = {case['variant']}")

        print("\n✅ 数据库 mock 数据灌入完成（幂等，重复执行安全）。")
        print(f"   后续联调可直接引用：exam_id={EXAM_ID}")
        for case in STUDENT_CASES:
            print(
                f"   {case['username']}: submission_id={case['submission_id']}"
            )
    finally:
        await conn.close()


ANSWER_VARIANTS = {
    "strong": {
        1: "B", 2: "B", 3: "B", 4: "B", 5: "A",
        6: "A", 7: "C", 8: "A", 9: "ABD", 10: "ABD",
        11: "AB", 12: "错", 13: "错", 14: "错",
        15: "列表可变、元组不可变；列表用方括号，元组用圆括号。",
        16: "递归是函数调用自身。必须包含基线条件和递归条件。",
        17: "__init__ 是构造方法，在创建对象时自动调用，用于初始化对象属性。",
        18: "def fibonacci(n):\n    if n < 2:\n        return n\n    a, b = 0, 1\n    for _ in range(n - 1):\n        a, b = b, a + b\n    return b",
        19: "def char_count(text):\n    result = {}\n    for ch in text:\n        result[ch] = result.get(ch, 0) + 1\n    return result",
        20: "def is_palindrome(text):\n    text = ''.join(ch.lower() for ch in text if ch.isalnum())\n    return text == text[::-1]",
    },
    "mixed": {
        1: "B", 2: "B", 3: "A", 4: "B", 5: "A",
        6: "A", 7: "C", 8: "C", 9: "ABD", 10: "AB",
        11: "AB", 12: "对", 13: "错", 14: "错",
        15: "列表可以修改，元组不能修改。",
        16: "递归是函数自己调用自己。需要终止条件和递归条件。",
        17: "用来创建对象。",
        18: "def fibonacci(n):\n    if n == 0:\n        return 0\n    if n == 1:\n        return 1\n    return fibonacci(n - 1) + fibonacci(n - 2)",
        19: "def char_count(text):\n    result = {}\n    for ch in text:\n        result[ch] = result.get(ch, 0) + 1\n    return result",
        20: "def is_palindrome(text):\n    text = text.replace(' ', '').lower()\n    return text == text[::-1]",
    },
    "weak": {
        1: "A", 2: "A", 3: "A", 4: "A", 5: "B",
        6: "B", 7: "D", 8: "C", 9: "ABC", 10: "A",
        11: "A", 12: "对", 13: "对", 14: "对",
        15: "列表和元组差不多。",
        16: "递归就是循环。",
        17: "不知道。",
        18: "def fibonacci(n):\n    return n",
        19: "def char_count(text):\n    return len(text)",
        20: "def is_palindrome(text):\n    return True",
    },
}


def make_student_word():
    """为每个学生生成不同的 20 题作答 Word 文件。"""
    from docx import Document

    os.makedirs("scripts/mock_files", exist_ok=True)

    for case in STUDENT_CASES:
        output_path = (
            f"scripts/mock_files/{case['username']}_python_exam.docx"
        )
        doc = Document()
        answers = ANSWER_VARIANTS[case["variant"]]

        for no in range(1, 21):
            doc.add_paragraph(f"第{no}题")
            answer = answers[no]
            if no in (18, 19, 20):
                doc.add_paragraph("```python")
                for line in answer.splitlines():
                    doc.add_paragraph(line)
                doc.add_paragraph("```")
            else:
                doc.add_paragraph(f"答：{answer}")

        try:
            doc.save(output_path)
            print(f"\n📝 学生作答 Word 已生成：{output_path}")
        except PermissionError:
            fallback_path = output_path.replace(".docx", "_new.docx")
            doc.save(fallback_path)
            print(
                f"\n⚠️  {output_path} 被占用，已改生成：{fallback_path}"
            )


if __name__ == "__main__":
    asyncio.run(seed())
    make_student_word()
